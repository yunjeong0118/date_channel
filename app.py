#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
채널별 편성표 → 요일별 PDF 생성기 (웹앱 버전, Streamlit)
=========================================================

Streamlit Community Cloud에 배포하면, 팀원은 exe/Excel 설치 없이
브라우저에서 링크만 열어 파일을 업로드하고 PDF/JPG를 내려받을 수 있습니다.

PDF 변환은 LibreOffice(soffice)를, PDF→JPG 변환은 poppler(pdftoppm)를
사용합니다 (Windows Excel 불필요). Streamlit Cloud에서 쓰려면 이 저장소에
packages.txt 로 `libreoffice`, `fonts-nanum`, `poppler-utils` 를 apt
패키지로 지정해둬야 합니다 (같이 준비해뒀습니다).
"""

import io
import shutil
import subprocess
import sys
import tempfile
import zipfile
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
import streamlit as st
from openpyxl.utils import get_column_letter

# ----------------------------------------------------------------------
# 요일 상수 / 핵심 로직 (channel_schedule_pdf.py 와 동일한 규칙)
# ----------------------------------------------------------------------
WEEKDAY_KR = ["월", "화", "수", "목", "금", "토", "일"]
WEEKDAY_EN = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
DAY_COL_CANDIDATES = [4, 6, 8, 10, 12, 14, 16]
HEADER_ROW = 6
DATA_FIRST_ROW = 6
DATA_LAST_ROW = 201
ROW_SHIFT = 2
MIRROR_SRC_COLS = (18, 19)

CHANNEL_PRIORITY = [
    ("PRIME2", "PRIME2"),
    ("PRIME+", "PRIME+"),
    ("PRIME", "PRIME"),
    ("SPOTV2", "SPOTV2"),
    ("PLUS", "SPOTV PLUS"),
    ("GOLF", "GOLF+"),
    ("SPOTV", "SPOTV"),
]
DISPLAY_ORDER_PRIORITY = ["SPOTV", "SPOTV2", "PRIME", "PRIME+", "PRIME2",
                           "SPOTV PLUS", "GOLF+"]


def detect_channel_name(sheet_title: str, filename: str) -> str:
    text = f"{sheet_title} {filename}".upper()
    for keyword, canonical in CHANNEL_PRIORITY:
        if keyword in text:
            return canonical
    return Path(filename).stem


def copy_cell(src_cell, dst_cell):
    dst_cell.value = src_cell.value
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)


def load_channel_files(file_paths, warn):
    channels = []
    for p in file_paths:
        wb = openpyxl.load_workbook(p)
        ws = wb.worksheets[0]
        name = detect_channel_name(ws.title, Path(p).name)
        channels.append({"name": name, "path": p, "ws": ws})
    for ch in channels:
        ws = ch["ws"]
        if ws.cell(row=HEADER_ROW, column=2).value != "시" or \
           ws.cell(row=HEADER_ROW, column=3).value != "분":
            warn(f"'{Path(ch['path']).name}' 파일이 예상 양식과 달라요. "
                 f"결과가 어긋날 수 있어요.")
    return channels


def order_channels(channels):
    def sort_key(ch):
        try:
            return (0, DISPLAY_ORDER_PRIORITY.index(ch["name"]))
        except ValueError:
            return (1, 0)
    return sorted(channels, key=sort_key)


def build_weekday_column_map(ref_ws):
    mapping = {}
    for col in DAY_COL_CANDIDATES:
        val = ref_ws.cell(row=HEADER_ROW, column=col).value
        if isinstance(val, datetime):
            mapping[val.weekday()] = col
    return mapping


def build_day_sheet(wb_out, sheet_name, target_date, day_src_col, channels, base_ws):
    ws_out = wb_out.create_sheet(title=sheet_name)
    n_ch = len(channels)

    ws_out.column_dimensions['A'].width = 1.25
    ws_out.column_dimensions['B'].width = 2.375
    ws_out.column_dimensions['C'].width = 2.0
    for i in range(n_ch):
        c1 = 4 + 2 * i
        ws_out.column_dimensions[get_column_letter(c1)].width = 6.125
        ws_out.column_dimensions[get_column_letter(c1 + 1)].width = 6.125
    mirror_col = 4 + 2 * n_ch
    ws_out.column_dimensions[get_column_letter(mirror_col)].width = 2.0
    ws_out.column_dimensions[get_column_letter(mirror_col + 1)].width = 2.375

    ws_out.row_dimensions[1].height = 4.5
    ws_out.row_dimensions[2].height = 22.35
    ws_out.row_dimensions[3].height = 5.25

    name_style_src = base_ws.cell(row=2, column=4)
    for i, ch in enumerate(channels):
        c1 = 4 + 2 * i
        c2 = c1 + 1
        ws_out.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
        cell = ws_out.cell(row=2, column=c1)
        cell.value = ch["name"]
        if name_style_src.has_style:
            cell.font = copy(name_style_src.font)
            cell.fill = copy(name_style_src.fill)
            cell.alignment = copy(name_style_src.alignment)
            cell.border = copy(name_style_src.border)
        cell.number_format = '@'

    for src_row in range(DATA_FIRST_ROW, DATA_LAST_ROW + 1):
        dest_row = src_row - ROW_SHIFT
        copy_cell(base_ws.cell(row=src_row, column=2), ws_out.cell(row=dest_row, column=2))
        copy_cell(base_ws.cell(row=src_row, column=3), ws_out.cell(row=dest_row, column=3))
        copy_cell(base_ws.cell(row=src_row, column=MIRROR_SRC_COLS[0]),
                  ws_out.cell(row=dest_row, column=mirror_col))
        copy_cell(base_ws.cell(row=src_row, column=MIRROR_SRC_COLS[1]),
                  ws_out.cell(row=dest_row, column=mirror_col + 1))
        rd = base_ws.row_dimensions.get(src_row)
        if rd and rd.height:
            ws_out.row_dimensions[dest_row].height = rd.height

    for i, ch in enumerate(channels):
        ws_src = ch["ws"]
        c1 = 4 + 2 * i
        c2 = c1 + 1
        for src_row in range(DATA_FIRST_ROW, DATA_LAST_ROW + 1):
            dest_row = src_row - ROW_SHIFT
            copy_cell(ws_src.cell(row=src_row, column=day_src_col),
                      ws_out.cell(row=dest_row, column=c1))
            copy_cell(ws_src.cell(row=src_row, column=day_src_col + 1),
                      ws_out.cell(row=dest_row, column=c2))

    for mc in base_ws.merged_cells.ranges:
        if mc.min_row < DATA_FIRST_ROW or mc.min_row > DATA_LAST_ROW:
            continue
        dr1, dr2 = mc.min_row - ROW_SHIFT, mc.max_row - ROW_SHIFT
        if mc.min_col == mc.max_col == 2:
            ws_out.merge_cells(start_row=dr1, start_column=2, end_row=dr2, end_column=2)
        elif mc.min_col == mc.max_col == 3:
            ws_out.merge_cells(start_row=dr1, start_column=3, end_row=dr2, end_column=3)
        elif mc.min_col == mc.max_col == MIRROR_SRC_COLS[0]:
            ws_out.merge_cells(start_row=dr1, start_column=mirror_col, end_row=dr2, end_column=mirror_col)
        elif mc.min_col == mc.max_col == MIRROR_SRC_COLS[1]:
            ws_out.merge_cells(start_row=dr1, start_column=mirror_col + 1, end_row=dr2, end_column=mirror_col + 1)

    for i, ch in enumerate(channels):
        ws_src = ch["ws"]
        c1 = 4 + 2 * i
        for mc in ws_src.merged_cells.ranges:
            if mc.min_row < DATA_FIRST_ROW or mc.min_row > DATA_LAST_ROW:
                continue
            if mc.min_col == day_src_col and mc.max_col == day_src_col + 1:
                dr1, dr2 = mc.min_row - ROW_SHIFT, mc.max_row - ROW_SHIFT
                ws_out.merge_cells(start_row=dr1, start_column=c1, end_row=dr2, end_column=c1 + 1)

    for i in range(n_ch):
        c1 = 4 + 2 * i
        cell = ws_out.cell(row=4, column=c1)
        cell.value = f"{target_date.month}월 {target_date.day}일 ({WEEKDAY_EN[target_date.weekday()]})"
        cell.number_format = '@'

    ws_out.sheet_view.zoomScale = 130
    ws_out.sheet_format.defaultRowHeight = 16.5

    ws_out.page_setup.orientation = 'landscape'
    ws_out.page_setup.paperSize = ws_out.PAPERSIZE_A3
    ws_out.page_setup.fitToWidth = 1
    ws_out.page_setup.fitToHeight = 0
    ws_out.sheet_properties.pageSetUpPr.fitToPage = True
    ws_out.print_area = f"A1:{get_column_letter(mirror_col + 1)}{ws_out.max_row}"
    ws_out.print_title_rows = '1:4'
    for m in ('left', 'right'):
        setattr(ws_out.page_margins, m, 0.2)
    for m in ('top', 'bottom'):
        setattr(ws_out.page_margins, m, 0.3)
    ws_out.page_margins.header = 0.1
    ws_out.page_margins.footer = 0.1

    return ws_out


def convert_to_pdf(xlsx_path: Path, outdir: Path) -> Path:
    soffice_bin = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice_bin:
        raise RuntimeError(
            "이 서버에 LibreOffice가 설치되어 있지 않아요. "
            "저장소의 packages.txt에 'libreoffice' 항목이 있는지 확인해주세요."
        )
    result = subprocess.run(
        [soffice_bin, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(xlsx_path)],
        capture_output=True, text=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(f"PDF 변환 실패: {result.stderr}")
    return outdir / (xlsx_path.stem + ".pdf")


def convert_pdf_to_jpgs(pdf_path: Path, outdir: Path, base_name: str):
    """PDF의 각 페이지를 JPG로 변환해 [(파일명, bytes), ...] 로 반환한다."""
    pdftoppm_bin = shutil.which("pdftoppm")
    if not pdftoppm_bin:
        raise RuntimeError(
            "이 서버에 poppler(pdftoppm)가 설치되어 있지 않아요. "
            "저장소의 packages.txt에 'poppler-utils' 항목이 있는지 확인해주세요."
        )
    prefix = outdir / base_name
    result = subprocess.run(
        [pdftoppm_bin, "-jpeg", "-r", "150", str(pdf_path), str(prefix)],
        capture_output=True, text=True, timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(f"JPG 변환 실패: {result.stderr}")
    jpg_files = sorted(outdir.glob(f"{base_name}-*.jpg")) or sorted(outdir.glob(f"{base_name}*.jpg"))
    if len(jpg_files) == 1:
        # 1페이지짜리는 페이지 번호 없이 깔끔한 이름으로
        return [(f"{base_name}.jpg", jpg_files[0].read_bytes())]
    return [(f.name, f.read_bytes()) for f in jpg_files]


def make_single_sheet_workbook(src_ws, sheet_name) -> openpyxl.Workbook:
    wb_single = openpyxl.Workbook()
    wb_single.remove(wb_single.active)
    new_ws = wb_single.create_sheet(title=sheet_name)
    for col, dim in src_ws.column_dimensions.items():
        if dim.width:
            new_ws.column_dimensions[col].width = dim.width
    for r, dim in src_ws.row_dimensions.items():
        if dim.height:
            new_ws.row_dimensions[r].height = dim.height
    for row in src_ws.iter_rows():
        for cell in row:
            nc = new_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                nc.font = copy(cell.font)
                nc.fill = copy(cell.fill)
                nc.border = copy(cell.border)
                nc.alignment = copy(cell.alignment)
                nc.number_format = cell.number_format
    for mc in src_ws.merged_cells.ranges:
        new_ws.merge_cells(str(mc))
    new_ws.page_setup = copy(src_ws.page_setup)
    new_ws.print_area = src_ws.print_area
    new_ws.print_title_rows = src_ws.print_title_rows
    new_ws.page_margins = copy(src_ws.page_margins)
    new_ws.sheet_properties.pageSetUpPr.fitToPage = True
    return wb_single


# ----------------------------------------------------------------------
# Streamlit UI
# ----------------------------------------------------------------------
st.set_page_config(page_title="채널별 편성표 → 요일별 PDF", page_icon="📺", layout="centered")
st.title("📺 채널별 편성표 → 요일별 PDF 생성기")
st.caption("채널별 주간 편성표 xlsx를 올리고 요일을 체크하면, 요일별로 모든 채널을 모은 PDF를 만들어줍니다.")

uploaded_files = st.file_uploader(
    "① 채널 편성표 파일 (xlsx, 여러 개 선택 가능)",
    type="xlsx",
    accept_multiple_files=True,
)

st.write("② 요일 선택")
cols = st.columns(7)
day_checked = []
for i, d in enumerate(WEEKDAY_KR):
    with cols[i]:
        day_checked.append(st.checkbox(d, value=(i >= 5), key=f"day_{i}"))

label = st.text_input("③ 파일명 라벨", value="채널별_편성표")

run = st.button("PDF 생성", type="primary", use_container_width=True)

if run:
    if not uploaded_files:
        st.warning("채널 편성표 파일을 먼저 업로드해주세요.")
        st.stop()
    selected_weekdays = [i for i, v in enumerate(day_checked) if v]
    if not selected_weekdays:
        st.warning("요일을 최소 1개 이상 체크해주세요.")
        st.stop()

    with st.spinner("생성 중..."):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmpdir = Path(tmpdir)
            local_paths = []
            for uf in uploaded_files:
                p = tmpdir / uf.name
                p.write_bytes(uf.getvalue())
                local_paths.append(p)

            warnings = []
            try:
                channels = load_channel_files(local_paths, warn=warnings.append)
                channels = order_channels(channels)
                base_ws = channels[0]["ws"]
                weekday_col_map = build_weekday_column_map(base_ws)

                missing = [WEEKDAY_KR[w] for w in selected_weekdays if w not in weekday_col_map]
                if missing:
                    st.error(f"기준 파일에서 다음 요일의 날짜를 찾지 못했습니다: {missing}")
                    st.stop()

                for w in warnings:
                    st.warning(w)
                st.info("인식된 채널: " + ", ".join(c["name"] for c in channels))

                wb_out = openpyxl.Workbook()
                wb_out.remove(wb_out.active)

                day_sheets = []
                for w in selected_weekdays:
                    col = weekday_col_map[w]
                    target_date = base_ws.cell(row=HEADER_ROW, column=col).value
                    sheet_name = f"{target_date.strftime('%m%d')}{WEEKDAY_KR[w]}요일"
                    build_day_sheet(wb_out, sheet_name, target_date, col, channels, base_ws)
                    day_sheets.append((sheet_name, target_date))

                combined_path = tmpdir / f"합본_{label}.xlsx"
                wb_out.save(combined_path)

                pdf_results = []
                jpg_results = []
                for sheet_name, target_date in day_sheets:
                    single_wb = make_single_sheet_workbook(wb_out[sheet_name], sheet_name)
                    single_path = tmpdir / f"_print_{sheet_name}.xlsx"
                    single_wb.save(single_path)
                    pdf_path = convert_to_pdf(single_path, tmpdir)

                    base_label = f"{target_date.strftime('%y%m%d')}({WEEKDAY_KR[target_date.weekday()]})_{label}"
                    final_pdf_name = f"{base_label}.pdf"
                    pdf_results.append((final_pdf_name, pdf_path.read_bytes()))

                    for jpg_name, jpg_bytes in convert_pdf_to_jpgs(pdf_path, tmpdir, base_label):
                        jpg_results.append((jpg_name, jpg_bytes))

                # ---- 전체를 zip 하나로 묶기 ----
                zip_buf = io.BytesIO()
                with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                    zf.writestr(combined_path.name, combined_path.read_bytes())
                    for fname, data in pdf_results:
                        zf.writestr(f"pdf/{fname}", data)
                    for fname, data in jpg_results:
                        zf.writestr(f"jpg/{fname}", data)
                zip_buf.seek(0)

                st.success("완료!")
                st.download_button(
                    "📦 전체 한 번에 다운로드 (zip: 합본 xlsx + PDF + JPG)",
                    data=zip_buf.getvalue(),
                    file_name=f"{label}.zip",
                    mime="application/zip",
                    type="primary",
                    use_container_width=True,
                )

                with st.expander("파일 하나씩 따로 받기"):
                    st.download_button(
                        "📊 합본 xlsx",
                        data=combined_path.read_bytes(),
                        file_name=combined_path.name,
                        use_container_width=True,
                    )
                    for fname, data in pdf_results:
                        st.download_button(
                            f"📄 {fname}",
                            data=data,
                            file_name=fname,
                            mime="application/pdf",
                            use_container_width=True,
                        )
                    for fname, data in jpg_results:
                        st.download_button(
                            f"🖼️ {fname}",
                            data=data,
                            file_name=fname,
                            mime="image/jpeg",
                            use_container_width=True,
                        )
            except Exception as e:
                st.error(f"오류가 발생했어요: {e}")
                raise
