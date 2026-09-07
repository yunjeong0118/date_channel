#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
채널별 편성표 취합 -> 요일별 PDF 생성기 (팀 배포용 GUI 버전)
================================================================

동료 PC에서 Python 설치 없이 exe로 실행하는 것을 전제로 만든 버전입니다.
PDF 변환은 LibreOffice 대신, PC에 이미 설치된 **Microsoft Excel**을
그대로 사용합니다 (win32com). 그래서 실행 PC에 반드시 Excel이
설치되어 있어야 합니다.

[사용 방법 - exe 버전]
    1) 프로그램 실행
    2) "채널 파일 추가" 로 채널별 주간 편성표 xlsx 여러 개 선택
    3) 원하는 요일 체크
    4) 저장 폴더 선택
    5) "PDF 생성" 클릭

[전제 조건 - 입력 파일 형식]
    - 각 파일은 SPOTV 주간 편성표 표준 양식(첫 시트 = 이번 주,
      6행 헤더: B6='시', C6='분', D/F/H/J/L/N/P 6행 = 월~일 날짜,
      7~201행 = 5분 단위 편성 데이터)이어야 합니다.
"""

import os
import re
import sys
import threading
import traceback
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


# ----------------------------------------------------------------------
# 요일 상수
# ----------------------------------------------------------------------
WEEKDAY_KR = ["월", "화", "수", "목", "금", "토", "일"]
WEEKDAY_EN = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
DAY_COL_CANDIDATES = [4, 6, 8, 10, 12, 14, 16]
HEADER_ROW = 6
DATA_FIRST_ROW = 6
DATA_LAST_ROW = 201
ROW_SHIFT = 2
MIRROR_SRC_COLS = (18, 19)

CHANNEL_PRIORITY = [
    ("PRIME2", "PRIME2"),
    ("PRIME+", "PRIME+"),
    ("PRIME", "PRIME"),
    ("SPOTV2", "SPOTV2"),
    ("PLUS", "SPOTV PLUS"),
    ("GOLF", "GOLF+"),
    ("SPOTV", "SPOTV"),
]
DISPLAY_ORDER_PRIORITY = ["SPOTV", "SPOTV2", "PRIME", "PRIME+", "PRIME2",
                           "SPOTV PLUS", "GOLF+"]


def detect_channel_name(sheet_title: str, filename: str) -> str:
    text = f"{sheet_title} {filename}".upper()
    for keyword, canonical in CHANNEL_PRIORITY:
        if keyword in text:
            return canonical
    return Path(filename).stem


def copy_cell(src_cell, dst_cell):
    dst_cell.value = src_cell.value
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)


def load_channel_files(file_paths, log=print):
    channels = []
    for p in file_paths:
        wb = openpyxl.load_workbook(p)
        ws = wb.worksheets[0]
        name = detect_channel_name(ws.title, Path(p).name)
        channels.append({"name": name, "path": p, "ws": ws})
    for ch in channels:
        ws = ch["ws"]
        if ws.cell(row=HEADER_ROW, column=2).value != "시" or \
           ws.cell(row=HEADER_ROW, column=3).value != "분":
            log(f"[경고] '{Path(ch['path']).name}' 파일이 예상 양식과 다릅니다. "
                f"결과가 어긋날 수 있어요.")
    return channels


def order_channels(channels):
    def sort_key(ch):
        try:
            return (0, DISPLAY_ORDER_PRIORITY.index(ch["name"]))
        except ValueError:
            return (1, 0)
    return sorted(channels, key=sort_key)


def build_weekday_column_map(ref_ws):
    mapping = {}
    for col in DAY_COL_CANDIDATES:
        val = ref_ws.cell(row=HEADER_ROW, column=col).value
        if isinstance(val, datetime):
            mapping[val.weekday()] = col
    return mapping


def build_day_sheet(wb_out, sheet_name, target_date, day_src_col, channels, base_ws):
    ws_out = wb_out.create_sheet(title=sheet_name)
    n_ch = len(channels)

    ws_out.column_dimensions['A'].width = 1.25
    ws_out.column_dimensions['B'].width = 2.375
    ws_out.column_dimensions['C'].width = 2.0
    for i in range(n_ch):
        c1 = 4 + 2 * i
        ws_out.column_dimensions[get_column_letter(c1)].width = 6.125
        ws_out.column_dimensions[get_column_letter(c1 + 1)].width = 6.125
    mirror_col = 4 + 2 * n_ch
    ws_out.column_dimensions[get_column_letter(mirror_col)].width = 2.0
    ws_out.column_dimensions[get_column_letter(mirror_col + 1)].width = 2.375

    ws_out.row_dimensions[1].height = 4.5
    ws_out.row_dimensions[2].height = 22.35
    ws_out.row_dimensions[3].height = 5.25

    name_style_src = base_ws.cell(row=2, column=4)
    for i, ch in enumerate(channels):
        c1 = 4 + 2 * i
        c2 = c1 + 1
        ws_out.merge_cells(start_row=2, start_column=c1, end_row=2, end_column=c2)
        cell = ws_out.cell(row=2, column=c1)
        cell.value = ch["name"]
        if name_style_src.has_style:
            cell.font = copy(name_style_src.font)
            cell.fill = copy(name_style_src.fill)
            cell.alignment = copy(name_style_src.alignment)
            cell.border = copy(name_style_src.border)
        cell.number_format = '@'

    for src_row in range(DATA_FIRST_ROW, DATA_LAST_ROW + 1):
        dest_row = src_row - ROW_SHIFT
        copy_cell(base_ws.cell(row=src_row, column=2), ws_out.cell(row=dest_row, column=2))
        copy_cell(base_ws.cell(row=src_row, column=3), ws_out.cell(row=dest_row, column=3))
        copy_cell(base_ws.cell(row=src_row, column=MIRROR_SRC_COLS[0]),
                  ws_out.cell(row=dest_row, column=mirror_col))
        copy_cell(base_ws.cell(row=src_row, column=MIRROR_SRC_COLS[1]),
                  ws_out.cell(row=dest_row, column=mirror_col + 1))
        rd = base_ws.row_dimensions.get(src_row)
        if rd and rd.height:
            ws_out.row_dimensions[dest_row].height = rd.height

    for i, ch in enumerate(channels):
        ws_src = ch["ws"]
        c1 = 4 + 2 * i
        c2 = c1 + 1
        for src_row in range(DATA_FIRST_ROW, DATA_LAST_ROW + 1):
            dest_row = src_row - ROW_SHIFT
            copy_cell(ws_src.cell(row=src_row, column=day_src_col),
                      ws_out.cell(row=dest_row, column=c1))
            copy_cell(ws_src.cell(row=src_row, column=day_src_col + 1),
                      ws_out.cell(row=dest_row, column=c2))

    for mc in base_ws.merged_cells.ranges:
        if mc.min_row < DATA_FIRST_ROW or mc.min_row > DATA_LAST_ROW:
            continue
        dr1, dr2 = mc.min_row - ROW_SHIFT, mc.max_row - ROW_SHIFT
        if mc.min_col == mc.max_col == 2:
            ws_out.merge_cells(start_row=dr1, start_column=2, end_row=dr2, end_column=2)
        elif mc.min_col == mc.max_col == 3:
            ws_out.merge_cells(start_row=dr1, start_column=3, end_row=dr2, end_column=3)
        elif mc.min_col == mc.max_col == MIRROR_SRC_COLS[0]:
            ws_out.merge_cells(start_row=dr1, start_column=mirror_col, end_row=dr2, end_column=mirror_col)
        elif mc.min_col == mc.max_col == MIRROR_SRC_COLS[1]:
            ws_out.merge_cells(start_row=dr1, start_column=mirror_col + 1, end_row=dr2, end_column=mirror_col + 1)

    for i, ch in enumerate(channels):
        ws_src = ch["ws"]
        c1 = 4 + 2 * i
        for mc in ws_src.merged_cells.ranges:
            if mc.min_row < DATA_FIRST_ROW or mc.min_row > DATA_LAST_ROW:
                continue
            if mc.min_col == day_src_col and mc.max_col == day_src_col + 1:
                dr1, dr2 = mc.min_row - ROW_SHIFT, mc.max_row - ROW_SHIFT
                ws_out.merge_cells(start_row=dr1, start_column=c1, end_row=dr2, end_column=c1 + 1)

    for i in range(n_ch):
        c1 = 4 + 2 * i
        cell = ws_out.cell(row=4, column=c1)
        cell.value = f"{target_date.month}월 {target_date.day}일 ({WEEKDAY_EN[target_date.weekday()]})"
        cell.number_format = '@'

    ws_out.sheet_view.zoomScale = 130
    ws_out.sheet_format.defaultRowHeight = 16.5

    ws_out.page_setup.orientation = 'landscape'
    ws_out.page_setup.paperSize = ws_out.PAPERSIZE_A3
    ws_out.page_setup.fitToWidth = 1
    ws_out.page_setup.fitToHeight = 0
    ws_out.sheet_properties.pageSetUpPr.fitToPage = True
    ws_out.print_area = f"A1:{get_column_letter(mirror_col + 1)}{ws_out.max_row}"
    ws_out.print_title_rows = '1:4'
    for m in ('left', 'right'):
        setattr(ws_out.page_margins, m, 0.2)
    for m in ('top', 'bottom'):
        setattr(ws_out.page_margins, m, 0.3)
    ws_out.page_margins.header = 0.1
    ws_out.page_margins.footer = 0.1

    return ws_out


def parse_days_selection(selected_flags):
    """selected_flags: dict {0..6: bool} -> ordered list of weekday indices"""
    return [i for i in range(7) if selected_flags.get(i)]


# ----------------------------------------------------------------------
# Excel COM 을 이용한 PDF 변환 (Windows + Microsoft Excel 필요)
# ----------------------------------------------------------------------
def export_workbook_sheet_to_pdf(xlsx_path: Path, sheet_name: str, pdf_path: Path, excel_app):
    wb = excel_app.Workbooks.Open(str(xlsx_path))
    try:
        ws = wb.Worksheets(sheet_name)
        ws.Select()
        # 0 = xlTypePDF
        wb.ActiveSheet.ExportAsFixedFormat(0, str(pdf_path))
    finally:
        wb.Close(SaveChanges=False)


def run_pipeline(file_paths, selected_weekdays, outdir: Path, label: str, log=print):
    outdir.mkdir(parents=True, exist_ok=True)

    channels = load_channel_files(file_paths, log=log)
    channels = order_channels(channels)
    log("인식된 채널(표시 순서): " + ", ".join(c["name"] for c in channels))

    base_ws = channels[0]["ws"]
    weekday_col_map = build_weekday_column_map(base_ws)

    missing = [WEEKDAY_KR[w] for w in selected_weekdays if w not in weekday_col_map]
    if missing:
        raise RuntimeError(f"기준 파일에서 다음 요일의 날짜를 찾지 못했습니다: {missing}")

    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    day_sheets = []
    for w in selected_weekdays:
        col = weekday_col_map[w]
        target_date = base_ws.cell(row=HEADER_ROW, column=col).value
        sheet_name = f"{target_date.strftime('%m%d')}{WEEKDAY_KR[w]}요일"
        build_day_sheet(wb_out, sheet_name, target_date, col, channels, base_ws)
        day_sheets.append((sheet_name, target_date))
        log(f"  - {sheet_name} 시트 생성 완료")

    combined_xlsx = outdir / f"합본_{label}.xlsx"
    wb_out.save(combined_xlsx)
    log(f"저장: {combined_xlsx}")

    # ---- 요일별 단일 시트 임시 파일 생성 ----
    tmp_paths = []
    for sheet_name, target_date in day_sheets:
        tmp_xlsx = outdir / f"_print_{sheet_name}.xlsx"
        wb_single = openpyxl.Workbook()
        wb_single.remove(wb_single.active)
        src_ws = wb_out[sheet_name]
        new_ws = wb_single.create_sheet(title=sheet_name)
        for col, dim in src_ws.column_dimensions.items():
            if dim.width:
                new_ws.column_dimensions[col].width = dim.width
        for r, dim in src_ws.row_dimensions.items():
            if dim.height:
                new_ws.row_dimensions[r].height = dim.height
        for row in src_ws.iter_rows():
            for cell in row:
                nc = new_ws.cell(row=cell.row, column=cell.column, value=cell.value)
                if cell.has_style:
                    nc.font = copy(cell.font)
                    nc.fill = copy(cell.fill)
                    nc.border = copy(cell.border)
                    nc.alignment = copy(cell.alignment)
                    nc.number_format = cell.number_format
        for mc in src_ws.merged_cells.ranges:
            new_ws.merge_cells(str(mc))
        new_ws.page_setup = copy(src_ws.page_setup)
        new_ws.print_area = src_ws.print_area
        new_ws.print_title_rows = src_ws.print_title_rows
        new_ws.page_margins = copy(src_ws.page_margins)
        new_ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb_single.save(tmp_xlsx)
        tmp_paths.append((tmp_xlsx, sheet_name, target_date))

    # ---- Excel COM 실행 (한 번만 띄워서 순차 변환) ----
    try:
        import win32com.client as win32
    except ImportError as e:
        raise RuntimeError(
            "pywin32 모듈이 없습니다. 이 exe는 Excel이 설치된 Windows PC에서만 동작합니다."
        ) from e

    log("Excel 실행 중... (화면에 잠깐 깜빡일 수 있어요)")
    excel_app = win32.gencache.EnsureDispatch('Excel.Application')
    excel_app.Visible = False
    excel_app.DisplayAlerts = False
    try:
        for tmp_xlsx, sheet_name, target_date in tmp_paths:
            final_name = f"{target_date.strftime('%y%m%d')}({WEEKDAY_KR[target_date.weekday()]})_{label}.pdf"
            final_path = outdir / final_name
            export_workbook_sheet_to_pdf(tmp_xlsx, sheet_name, final_path, excel_app)
            log(f"저장: {final_path}")
    finally:
        excel_app.Quit()
        for tmp_xlsx, _, _ in tmp_paths:
            try:
                tmp_xlsx.unlink()
            except OSError:
                pass

    log("완료.")


# ----------------------------------------------------------------------
# GUI
# ----------------------------------------------------------------------
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("채널별 편성표 → 요일별 PDF 생성기")
        self.geometry("640x520")
        self.resizable(False, False)

        self.file_paths = []

        pad = {"padx": 10, "pady": 6}

        frm_files = ttk.LabelFrame(self, text="① 채널 편성표 파일")
        frm_files.pack(fill="x", **pad)
        self.listbox = tk.Listbox(frm_files, height=6)
        self.listbox.pack(side="left", fill="both", expand=True, padx=(8, 0), pady=8)
        btns = ttk.Frame(frm_files)
        btns.pack(side="left", padx=8, pady=8)
        ttk.Button(btns, text="파일 추가", command=self.add_files).pack(fill="x", pady=2)
        ttk.Button(btns, text="선택 삭제", command=self.remove_selected).pack(fill="x", pady=2)
        ttk.Button(btns, text="모두 지우기", command=self.clear_files).pack(fill="x", pady=2)

        frm_days = ttk.LabelFrame(self, text="② 요일 선택")
        frm_days.pack(fill="x", **pad)
        self.day_vars = []
        for i, d in enumerate(WEEKDAY_KR):
            var = tk.BooleanVar(value=(i >= 5))  # 기본값: 토, 일 체크
            chk = ttk.Checkbutton(frm_days, text=d, variable=var)
            chk.grid(row=0, column=i, padx=10, pady=8)
            self.day_vars.append(var)

        frm_out = ttk.LabelFrame(self, text="③ 저장 폴더 / 파일 라벨")
        frm_out.pack(fill="x", **pad)
        self.outdir_var = tk.StringVar(value=str(Path.home() / "Desktop"))
        ttk.Entry(frm_out, textvariable=self.outdir_var, width=55).grid(row=0, column=0, padx=8, pady=6)
        ttk.Button(frm_out, text="찾아보기", command=self.choose_outdir).grid(row=0, column=1, padx=4)
        self.label_var = tk.StringVar(value="채널별_편성표")
        ttk.Label(frm_out, text="파일명 라벨:").grid(row=1, column=0, sticky="w", padx=8)
        ttk.Entry(frm_out, textvariable=self.label_var, width=30).grid(row=2, column=0, sticky="w", padx=8, pady=(0, 6))

        self.run_btn = ttk.Button(self, text="PDF 생성", command=self.on_run)
        self.run_btn.pack(pady=10)

        frm_log = ttk.LabelFrame(self, text="진행 상황")
        frm_log.pack(fill="both", expand=True, **pad)
        self.log_text = tk.Text(frm_log, height=10, state="disabled")
        self.log_text.pack(fill="both", expand=True, padx=6, pady=6)

    def log(self, msg):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", msg + "\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")
        self.update_idletasks()

    def add_files(self):
        paths = filedialog.askopenfilenames(
            title="채널별 편성표 xlsx 선택 (여러 개 선택 가능)",
            filetypes=[("Excel 파일", "*.xlsx")],
        )
        for p in paths:
            if p not in self.file_paths:
                self.file_paths.append(p)
                self.listbox.insert("end", Path(p).name)

    def remove_selected(self):
        for idx in reversed(self.listbox.curselection()):
            self.listbox.delete(idx)
            del self.file_paths[idx]

    def clear_files(self):
        self.listbox.delete(0, "end")
        self.file_paths = []

    def choose_outdir(self):
        d = filedialog.askdirectory(title="저장 폴더 선택")
        if d:
            self.outdir_var.set(d)

    def on_run(self):
        if not self.file_paths:
            messagebox.showwarning("알림", "채널 편성표 파일을 먼저 추가해 주세요.")
            return
        selected = [i for i, v in enumerate(self.day_vars) if v.get()]
        if not selected:
            messagebox.showwarning("알림", "요일을 최소 1개 이상 체크해 주세요.")
            return

        self.run_btn.configure(state="disabled")
        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")

        outdir = Path(self.outdir_var.get())
        label = self.label_var.get().strip() or "채널별_편성표"
        files = list(self.file_paths)

        def worker():
            try:
                run_pipeline(files, selected, outdir, label, log=self.log)
                self.after(0, lambda: messagebox.showinfo("완료", f"PDF 생성이 끝났어요.\n저장 위치: {outdir}"))
                self.after(0, lambda: os.startfile(outdir))
            except Exception as e:
                tb = traceback.format_exc()
                self.log("오류 발생:\n" + tb)
                self.after(0, lambda: messagebox.showerror("오류", str(e)))
            finally:
                self.after(0, lambda: self.run_btn.configure(state="normal"))

        threading.Thread(target=worker, daemon=True).start()


if __name__ == "__main__":
    if sys.platform != "win32":
        print("이 GUI 버전은 Windows + Microsoft Excel 환경에서 실행해야 합니다.")
    app = App()
    app.mainloop()
